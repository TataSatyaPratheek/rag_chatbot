# tests/conftest.py
import os
import pytest
import tempfile
from pathlib import Path
import fitz  # PyMuPDF
from pptx import Presentation
import numpy as np
from unittest.mock import MagicMock # Added import
from PIL import Image

from mmrag.document_processing import (
    PDFProcessor, PowerPointProcessor, TableDetector, 
    EnhancedVisualProcessor
)
from mmrag.vectordb import ChromaStore
from mmrag.llm import OllamaClient, ContentUnderstanding # Keep OllamaClient import

@pytest.fixture
def test_data_dir():
    """Path to the test data directory."""
    return Path(__file__).parent / "data"

@pytest.fixture
def sample_data_dir(test_data_dir):
    """Path to the sample data directory."""
    return test_data_dir / "samples"

@pytest.fixture
def large_data_dir(test_data_dir):
    """Path to the large data directory."""
    return test_data_dir / "large"

@pytest.fixture
def temp_dir():
    """Provide a temporary directory that's automatically cleaned up."""
    with tempfile.TemporaryDirectory() as temp_dir:
        yield Path(temp_dir)

@pytest.fixture
def sample_pdf_path(sample_data_dir):
    """Path to a sample PDF file."""
    pdf_path = sample_data_dir / "pdf" / "sample.pdf"
    if not pdf_path.exists():
        # Create a simple PDF for testing if it doesn't exist
        pdf_path.parent.mkdir(parents=True, exist_ok=True)
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 50), "Sample PDF for testing")
        page.insert_text((50, 70), "This is a second line of text")
        doc.save(pdf_path)
    return pdf_path

@pytest.fixture
def sample_ppt_path(sample_data_dir):
    """Path to a sample PowerPoint file."""
    ppt_path = sample_data_dir / "ppt" / "sample.pptx"
    if not ppt_path.exists():
        # Create a simple PPTX for testing if it doesn't exist
        ppt_path.parent.mkdir(parents=True, exist_ok=True)
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Sample Presentation"
        prs.save(ppt_path)
    return ppt_path

@pytest.fixture
def sample_table_image(sample_data_dir):
    """Create a sample table image for testing."""
    table_path = sample_data_dir / "tables" / "sample_table.png"
    if not table_path.exists():
        table_path.parent.mkdir(parents=True, exist_ok=True)
        # Create a simple table image
        img = Image.new('RGB', (300, 200), color='white')
        # Draw table lines
        table_array = np.ones((200, 300, 3), dtype=np.uint8) * 255
        # Horizontal lines
        table_array[0:2, :] = 0  # Top border
        table_array[50:52, :] = 0  # Header row
        table_array[100:102, :] = 0  # Middle row
        table_array[150:152, :] = 0  # Middle row
        table_array[198:200, :] = 0  # Bottom border
        # Vertical lines
        table_array[:, 0:2] = 0  # Left border
        table_array[:, 99:101] = 0  # Left column
        table_array[:, 199:201] = 0  # Middle column
        table_array[:, 298:300] = 0  # Right border
        # Convert to PIL Image and save
        table_img = Image.fromarray(table_array.astype('uint8'))
        table_img.save(table_path)
    return table_path

@pytest.fixture
def sample_chart_image(sample_data_dir):
    """Create a sample chart image for testing."""
    chart_path = sample_data_dir / "charts" / "sample_barchart.png"
    if not chart_path.exists():
        chart_path.parent.mkdir(parents=True, exist_ok=True)
        # Create a simple bar chart image
        img = Image.new('RGB', (300, 200), color='white')
        chart_array = np.ones((200, 300, 3), dtype=np.uint8) * 255
        # Draw axes
        chart_array[150:152, 30:270] = 0  # X-axis
        chart_array[30:151, 30:32] = 0  # Y-axis
        # Draw bars
        chart_array[100:150, 50:70] = [200, 0, 0]  # Bar 1
        chart_array[70:150, 100:120] = [0, 200, 0]  # Bar 2
        chart_array[120:150, 150:170] = [0, 0, 200]  # Bar 3
        chart_array[50:150, 200:220] = [200, 200, 0]  # Bar 4
        # Convert to PIL Image and save
        chart_img = Image.fromarray(chart_array.astype('uint8'))
        chart_img.save(chart_path)
    return chart_path

@pytest.fixture
def pdf_processor():
    """PDF processor instance for testing."""
    return PDFProcessor(extract_tables=True, extract_images=True)

@pytest.fixture
def ppt_processor():
    """PowerPoint processor instance for testing."""
    return PowerPointProcessor(extract_tables=True, extract_images=True)

@pytest.fixture
def table_detector():
    """Table detector instance for testing."""
    return TableDetector()

@pytest.fixture
def visual_processor():
    """Enhanced visual processor instance for testing."""
    return EnhancedVisualProcessor(detect_charts=True)

@pytest.fixture
def vector_store(temp_dir):
    """ChromaDB vector store for testing."""
    return ChromaStore(
        persist_directory=temp_dir / "chroma",
        collection_name="test_collection"
    )

@pytest.fixture
def mock_ollama_client(monkeypatch):
    """Mock Ollama client that returns predefined responses."""
    mock_client = OllamaClient()
    
    # Use MagicMock for the methods to allow assert_called_once etc.
    mock_generate = MagicMock(return_value="This is a mock LLM response for testing purposes.")
    mock_chat = MagicMock(return_value="This is a mock chat response for testing purposes.")
    
    # Patch the methods on the instance
    monkeypatch.setattr(mock_client, "generate_sync", mock_generate)
    monkeypatch.setattr(mock_client, "chat_sync", mock_chat)
    
    return mock_client

@pytest.fixture
def content_analyzer(mock_ollama_client):
    """Content understanding analyzer with mocked LLM client."""
    return ContentUnderstanding(llm_client=mock_ollama_client)

# Create a sample large document for stress testing
@pytest.fixture
def large_pdf_path(large_data_dir):
    """Create a large PDF document for stress testing."""
    pdf_path = large_data_dir / "large_sample.pdf"
    if not pdf_path.exists():
        pdf_path.parent.mkdir(parents=True, exist_ok=True)
        
        doc = fitz.open()
        # Create 100 pages with significant content
        for i in range(100):
            page = doc.new_page()
            for j in range(50):  # 50 text blocks per page
                y_pos = j * 15 + 50
                page.insert_text((50, y_pos), f"This is line {j+1} on page {i+1} for stress testing")
        
        doc.save(pdf_path)
    return pdf_path
